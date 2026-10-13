"""Turn many file types into plain text so the model can read them.

Nothing here knows about modules. It only extracts text. Images and video have no
readable text, so the agent gets a note saying so and has to decide another way.
"""
from __future__ import annotations

import json
from pathlib import Path

TEXT_EXT = {
    ".txt", ".md", ".csv", ".tsv", ".json", ".java", ".py", ".c", ".cpp", ".h", ".js", ".html",
    ".htm", ".xml", ".yaml", ".yml", ".log", ".tex", ".r", ".m", ".sql",
}
BIG = 15 * 1024 * 1024  # notebooks above this size are read as raw text, not parsed


def extract_text(path: Path, max_chars: int = 1500) -> str:
    ext = path.suffix.lower()
    size = path.stat().st_size
    try:
        if ext in TEXT_EXT:
            text = _head(path, max_chars)
        elif ext == ".ipynb":
            text = _head(path, max_chars) if size > BIG else _ipynb(path, max_chars)
        elif ext == ".pdf":
            text = _pdf(path, max_chars)
        elif ext == ".docx":
            text = _docx(path, max_chars)
        elif ext == ".pptx":
            text = _pptx(path, max_chars)
        elif ext in {".xlsx", ".xlsm"}:
            text = _xlsx(path, max_chars)
        else:
            return f"(no readable text, {ext or 'unknown'} file of {size} bytes)"
    except ImportError as e:
        return f"NOTE: cannot read {ext} files because the library {e.name} is missing. Use the file name only."
    except Exception as e:  # corrupt, encrypted or unusual files
        return f"NOTE: could not read this {ext} file ({type(e).__name__}). Use the file name only."
    text = text.strip()
    if not text:
        return f"(no extractable text in this {ext} file, it may be made of images)"
    if len(text) > max_chars:
        text = text[:max_chars] + "\n...[truncated]"
    return text


def _head(path: Path, n: int) -> str:
    with path.open(encoding="utf-8", errors="replace") as fh:
        return fh.read(n + 1)


def _ipynb(path: Path, n: int) -> str:
    nb = json.loads(path.read_text(encoding="utf-8", errors="replace"))
    parts, total = [], 0
    for cell in nb.get("cells", []):
        src = cell.get("source", "")
        src = "".join(src) if isinstance(src, list) else str(src)
        parts.append(f"[{cell.get('cell_type', 'cell')}] {src}")
        total += len(src)
        if total >= n:
            break
    return "\n".join(parts)


def _pdf(path: Path, n: int) -> str:
    from pypdf import PdfReader

    out, total = [], 0
    for page in PdfReader(str(path)).pages[:8]:
        t = page.extract_text() or ""
        out.append(t)
        total += len(t)
        if total >= n:
            break
    return "\n".join(out)


def _docx(path: Path, n: int) -> str:
    import docx

    d = docx.Document(str(path))
    parts, total = [], 0
    for p in d.paragraphs:
        if p.text.strip():
            parts.append(p.text)
            total += len(p.text)
            if total >= n:
                return "\n".join(parts)
    for table in d.tables[:3]:
        for row in table.rows[:10]:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)


def _pptx(path: Path, n: int) -> str:
    from pptx import Presentation

    parts, total = [], 0
    for i, slide in enumerate(Presentation(str(path)).slides, 1):
        texts = [s.text_frame.text for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()]
        if texts:
            parts.append(f"[slide {i}] " + " / ".join(texts))
            total += sum(len(t) for t in texts)
        if total >= n or i >= 20:
            break
    return "\n".join(parts)


def _xlsx(path: Path, n: int) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    parts, total = [], 0
    try:
        for ws in wb.worksheets[:3]:
            parts.append(f"[sheet {ws.title}]")
            for r, row in enumerate(ws.iter_rows(values_only=True)):
                line = " | ".join("" if c is None else str(c) for c in row).strip(" |")
                if line:
                    parts.append(line)
                    total += len(line)
                if r >= 25 or total >= n:
                    break
    finally:
        wb.close()
    return "\n".join(parts)
